"""Menyusun laporan ilmiah Kelompok 1 CAYMAN (Kasus 38) sebagai .docx.
Format sesuai panduan: single column, A4, spasi 1.15, margin 2.5 cm.
Semua angka berasal dari outputs/ (hasil eksekusi nyata di Colab & uji browser)."""
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
FIG = HERE / "fig"
OUT = HERE / "Tubes_DL_Kelompok01_Oil_Spill_SAR.docx"
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.15
pf.space_after = Pt(4)
pf.space_before = Pt(0)

for name, size in (("Heading 1", 12.5), ("Heading 2", 11.5)):
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
    st.font.color.rgb = RGBColor(0, 0, 0)
    rf = st.element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    st.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.keep_with_next = True


def add_runs(par, text, size=None, italic_all=False):
    """Markup sederhana: **tebal**, *miring*."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if italic_all:
            r.italic = True
        if size:
            r.font.size = Pt(size)
    return par


def P(text, align="justify", size=None, italic=False, space_after=None, indent=True):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent and align == "justify":
        p.paragraph_format.first_line_indent = Cm(0.75)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return add_runs(p, text, size, italic)


def B(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.9 + 0.6 * level)
    p.paragraph_format.space_after = Pt(2)
    return add_runs(p, text)


def H1(t): doc.add_heading(t, level=1)
def H2(t): doc.add_heading(t, level=2)


def caption(text, before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8 if not before else 3)
    p.paragraph_format.space_before = Pt(3 if not before else 6)
    if before:
        p.paragraph_format.keep_with_next = True
    label, rest = text.split(".", 1)
    r = p.add_run(label + "."); r.bold = True; r.font.size = Pt(9.5)
    add_runs(p, rest, size=9.5)


def figure(path, width_cm, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(text)


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for edge, val in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val.get("val", "single")); el.set(qn("w:sz"), str(val.get("sz", 6)))
        el.set(qn("w:color"), val.get("color", "000000")); el.set(qn("w:space"), "0")
        borders.append(el)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(title, header, rows, widths_cm, size=9, bold_rows=(), note=None, align_first_left=True):
    caption(title, before=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths_cm[j])
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and align_first_left) else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(c.paragraphs[0], str(val), size=size)
            if i == 0 or i - 1 in bold_rows:
                for r in c.paragraphs[0].runs: r.bold = True
            if i == 0:
                shade(c, "E7EEF7")
                set_cell_border(c, top={"sz": 10}, bottom={"sz": 6})
            if i == len(rows):
                set_cell_border(c, bottom={"sz": 10})
            if i - 1 in bold_rows:
                shade(c, "F2F7EC")
            if i < len(rows):  # jaga tabel tetap utuh satu halaman
                c.paragraphs[0].paragraph_format.keep_with_next = True
        trPr = t.rows[i]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    if note:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        add_runs(p, note, size=8.5)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_number_footer(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; r._r.append(el)
    r.font.size = Pt(9)
    hp = section.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Tugas Besar Deep Learning · Kelompok 1 CAYMAN · Kasus 38"); hr.font.size = Pt(8.5); hr.italic = True


# ================================== COVER =====================================
def cover_line(text, size=12, bold=False, before=0, after=0, italic=False):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    return p

cover_line("LAPORAN TUGAS BESAR", 16, True, before=24)
cover_line("DEEP LEARNING", 14, True, after=18)
cover_line("Perbandingan Machine Learning dan Deep Learning untuk Deteksi,", 13, True)
cover_line("Pemetaan Poligon, dan Perencanaan Respons Tumpahan Minyak", 13, True)
cover_line("pada Citra Sentinel-1 SAR di Browser", 13, True, after=4)
cover_line("(Kasus 38: Marine Oil Spill Detection and Mapping via Satellite SAR)", 11, after=22)
lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER; lp.paragraph_format.space_after = Pt(20)
lp.add_run().add_picture(str(FIG / "logo_upnvjt.png"), width=Cm(5.2))
cover_line("Disusun oleh:", 12, after=2)
cover_line("Kelompok 1 — CAYMAN", 12, True, after=6)
members = [("Ade Rizky Darmawan (Ketua)", "23083010080"), ("Arkananta Daniswara Handoyo", "23083010059"),
           ("Muhammad Arsyad Alzam", "23083010082"), ("Choirul Amin", "22083010050"),
           ("Hana Titania Sastrian", "23083010056"), ("Zaydan Arief Athallah", "23083010063")]
mt = doc.add_table(rows=len(members), cols=2); mt.alignment = WD_TABLE_ALIGNMENT.CENTER; mt.autofit = False
for i, (nm, npm) in enumerate(members):
    for j, (val, w, al) in enumerate(((nm, 7.0, WD_ALIGN_PARAGRAPH.LEFT), ("NPM : " + npm, 4.2, WD_ALIGN_PARAGRAPH.LEFT))):
        c = mt.cell(i, j); c.width = Cm(w); pp = c.paragraphs[0]; pp.alignment = al
        pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = 1.0
        pp.add_run(val).font.size = Pt(11.5)
cover_line("Dosen Pengampu:", 12, before=16)
cover_line("Dr. I Gede Susrama Mas Diyasa, ST., MT.", 12, after=6)
cover_line("Dosen Praktisi:", 12)
cover_line("Kahpi Baiquni Arifani, S.Kom., M.Kom.", 12, after=26)
for t in ("PROGRAM STUDI SAINS DATA", "FAKULTAS ILMU KOMPUTER", "UPN “VETERAN” JAWA TIMUR", "2026"):
    cover_line(t, 12, True)

# Section isi: nomor halaman mulai 1, header/footer hanya di sini
sec = doc.add_section(WD_SECTION.NEW_PAGE)
sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); sec._sectPr.append(pg)
page_number_footer(sec)

# ============================== JUDUL & IDENTITAS ==============================
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER; tp.paragraph_format.space_after = Pt(6)
r = tp.add_run("Perbandingan Machine Learning dan Deep Learning untuk Deteksi, Pemetaan Poligon, dan Perencanaan Respons Tumpahan Minyak pada Citra Sentinel-1 SAR di Browser")
r.bold = True; r.font.size = Pt(15)

P("Kelompok 1 — **CAYMAN** · Kasus 38: *Marine Oil Spill Detection and Mapping via Satellite SAR*", align="center", size=10.5, space_after=2)
P("Kode & dokumentasi: github.com/aderizkydarmawan/cayman-oil-spill-sar · Aplikasi web: cayman-kelompok1.vercel.app", align="center", size=9.5, space_after=10)

# ================================== ABSTRAK ===================================
ab = doc.add_paragraph(); ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ab.paragraph_format.left_indent = Cm(0.8); ab.paragraph_format.right_indent = Cm(0.8)
rr = ab.add_run("Abstrak — "); rr.bold = True; rr.font.size = Pt(10)
add_runs(ab, (
    "Tumpahan minyak dari fasilitas migas lepas pantai, pipa bawah laut, dan kapal tanker harus dipetakan secepat mungkin agar respons "
    "HSSE tepat sasaran. Penelitian ini membandingkan dua model *machine learning* (Random Forest dan LightGBM dengan 13 fitur tekstur "
    "buatan) dan dua model *deep learning* (U-Net Lite yang dilatih dari nol dan DeepLabV3+ MobileNetV2 pralatih) untuk segmentasi tumpahan "
    "minyak pada citra Sentinel-1 SAR dataset SOS (2.850 latih, 504 validasi, 839 uji). Semua model memakai split, aturan threshold "
    "(dipilih pada data validasi), dan metrik yang sama. Pada data uji, DeepLabV3+ mencapai Dice 0,867 dan IoU 0,765, disusul U-Net Lite "
    "(0,855), LightGBM (0,806), dan Random Forest (0,804), seluruhnya di atas baseline *dark-spot* (0,726). Model DL unggul 5–6 poin Dice "
    "terutama karena precision yang lebih tinggi, sedangkan model ML sangat kecil (0,24–1,8 MB). Kelima model diekspor ke ONNX (DL "
    "dikuantisasi INT8 tanpa penurunan Dice yang berarti) dan dijalankan di browser. Aplikasi web juga mengubah mask menjadi poligon per "
    "slick, mensimulasikan arah dan laju penyebaran (drift 3% angin + arus, penyebaran Fay dan difusi Okubo), memperkirakan panjang *oil "
    "boom*, serta menghitung efisiensi biaya dan penurunan risiko HSSE akibat berkurangnya patroli lapangan."), size=10)
kw = doc.add_paragraph(); kw.paragraph_format.left_indent = Cm(0.8); kw.paragraph_format.right_indent = Cm(0.8); kw.paragraph_format.space_after = Pt(8)
r = kw.add_run("Kata kunci: "); r.bold = True; r.font.size = Pt(10)
add_runs(kw, "oil spill, SAR, Sentinel-1, segmentasi semantik, random forest, LightGBM, U-Net, DeepLabV3+, edge AI, oil boom, HSSE.", size=10, italic_all=True)

# ================================ PENDAHULUAN =================================
H1("1. Pendahuluan")
P("Tumpahan minyak adalah salah satu ancaman terbesar bagi ekosistem laut, dan sebagian besar terkait aktivitas industri migas lepas "
  "pantai: semburan dari rig pengeboran, kecelakaan tanker, hingga buangan rutin kapal [6]. Pada pilar **HSSE** (*Health, Safety, Security "
  "and Environment*), tim respons membutuhkan lokasi, luas, dan arah gerak tumpahan secepat mungkin untuk menentukan penempatan *oil boom*, "
  "*skimmer*, dan kapal penanggulangan [1]. Pengecekan manual melalui patroli udara atau laut mahal dan memaparkan personel pada uap "
  "hidrokarbon serta risiko kecelakaan.")
P("Radar apertur sintetis (SAR) menjadi sensor utama pemantauan tumpahan minyak karena bekerja siang–malam dan menembus awan. Lapisan "
  "minyak meredam gelombang kapiler sehingga tampak sebagai area gelap pada citra SAR [1]. Liao dkk. menemukan lebih dari 80% kejadian "
  "tumpahan di Teluk Jiaozhou terjadi pada malam hari [8]. Tantangan utamanya adalah fenomena *look-alike* (angin lemah, lapisan biogenik) "
  "yang juga tampak gelap [2], [4], [7]. Tabel 1 merangkum penelitian rujukan; angka antarpenelitian umumnya **tidak dapat dibandingkan "
  "langsung** karena dataset dan protokolnya berbeda, kecuali Zhu dkk. [3] yang memperkenalkan dataset SOS yang juga kami gunakan.")
P("Kontribusi proyek ini: (1) *pipeline* yang diaudit, termasuk pemeriksaan kebocoran antar split; (2) perbandingan adil model ML dan "
  "DL pada protokol yang sama, termasuk uji lintas sensor; (3) aplikasi web dengan inferensi di sisi klien untuk kelima model, poligon per "
  "slick, dan GeoJSON; serta (4) simulasi arah dan laju penyebaran, estimasi kebutuhan *oil boom*, dan analisis biaya serta risiko HSSE.")
table("Tabel 1. Ringkasan penelitian rujukan (state of the art).",
      ["Rujukan", "Data", "Metode", "Hasil utama (sesuai paper)"],
      [["Zhu dkk., 2022 [3]", "SOS: ALOS PALSAR & Sentinel-1A", "CBD-Net (multiskala, scSE, supervisi tepi)", "Sentinel-1: mIoU 83,42%, F1 87,87%"],
       ["Shaban dkk., 2021 [5]", "SAR, dataset tidak seimbang", "CNN 23 lapis + U-Net 5 tahap", "Dice 80%, precision 84%"],
       ["Hasimoto-Beltran dkk., 2023 [6]", "16 citra ENVISAT-ASAR", "M-DNN multi-kanal (U-Net + ResNet)", "F1 hingga 98,24% (dataset sendiri)"],
       ["Li dkk., 2023 [7]", "Sentinel-1", "Algoritma DL *self-evolving*", "F1 0,8423 → 0,8896 (21 siklus)"],
       ["Liao dkk., 2023 [8]", "35 citra Sentinel-1 PolSAR", "DeepLabv3+ + fitur polarimetrik", "OA 0,9812; MIoU 0,9599"],
       ["Das dkk., 2024 [9]", "SAR, 5 kelas", "U-Net multi-backbone (EfficientNet-B3)", "mIoU 76,53%; IoU minyak 62,08%"],
       ["Petalas dkk., 2025 [10]", "Sentinel-1, Laut Aegea", "U-Net operasional + model dispersi", "IoU kelas minyak 0,30"]],
      [3.4, 3.6, 4.4, 4.6], size=8.5)
# ================================= METODOLOGI =================================
H1("2. Metodologi")
H2("2.1 Dataset, audit, dan split")
P("Dataset *Deep-SAR Oil Spill* (SOS) [3] diunduh dari Kaggle melalui kagglehub [13] (lisensi CC BY 4.0). SOS berisi 8.070 pasangan "
  "citra–mask PNG 256×256 dari ALOS PALSAR (Teluk Meksiko; 3.101 latih / 776 uji) dan Sentinel-1A VV (Teluk Persia; 3.354 latih / 839 "
  "uji). Audit menemukan: semua pasangan cocok tanpa NaN/Inf; ketiga kanal RGB identik sehingga model memakai 1 kanal; label latih memiliki "
  "2,2–2,5% piksel bernilai antara sehingga dibinerkan dengan aturan ≥128; tidak ada duplikat eksak, dan pemeriksaan *near-duplicate* "
  "dengan *average hash* terhadap 8 rotasi/flip hanya menemukan 9 citra uji yang mirip citra latih (dampaknya diuji di Subbab 3.1).")
P("Eksperimen memakai **Sentinel-1** (sampel terbanyak, label ambigu lebih sedikit, data terbuka). Data latih dibagi 2.850 latih dan 504 "
  "validasi (*seed* 42) dengan mengelompokkan komponen *near-duplicate* agar tidak terpisah; set uji 839 citra hanya dipakai untuk evaluasi "
  "akhir. Citra dinormalisasi *z-score* x' = (x/255 − 0,3834)/0,2020 dengan statistik dari data latih. Augmentasi (8 rotasi/flip identik "
  "untuk citra dan mask, jitter kecerahan ±10%) hanya pada data latih.")
figure(FIG / "pipeline.png", 15.0, "Gambar 1. Alur sistem, dari data hingga aplikasi web.")

H2("2.2 Model Machine Learning dan Deep Learning")
B("**Random Forest** [15] dan **LightGBM** [16] mengklasifikasi setiap piksel dari 13 fitur buatan tangan: intensitas; rata-rata lokal "
  "5/11/21/41 piksel; simpangan baku lokal 5/11/21 (tekstur); gradien Sobel (tepi); kontras rata-rata 5 vs 41; selisih terhadap rata-rata "
  "citra; serta minimum dan maksimum lokal 11 piksel. Fitur ditulis sebagai modul PyTorch agar identik di notebook dan browser. Model "
  "dilatih pada 300 piksel acak per citra latih (855.000 baris) dengan ukuran dibatasi agar ringan di web: RF 20 pohon kedalaman 12, "
  "LightGBM 100 pohon 31 daun.")
B("**U-Net Lite** adalah U-Net [11] 4 level dengan kanal 16–256 (1,94 juta parameter) yang dilatih dari nol. **DeepLabV3+** [17] memakai "
  "*encoder* MobileNetV2 [18] pralatih ImageNet (4,38 juta parameter; bobot kanal pertama dijumlahkan menjadi 1 kanal). Keduanya dilatih "
  "dengan resep yang sama: loss 0,5·BCE + 0,5·Dice, AdamW (lr 10⁻³, *weight decay* 10⁻⁴), jadwal *cosine* 40 epoch, *batch* 16, dan "
  "*mixed precision* pada GPU Tesla T4 (Google Colab). Checkpoint terbaik dipilih dari Dice validasi.")
P("Threshold setiap model dipilih melalui *sweep* 0,20–0,80 pada data **validasi**. Metrik utama adalah Dice (F1 piksel), IoU, precision, "
  "recall, dan PR-AUC dengan agregasi global (TP/FP/FN/TN dijumlahkan dari semua piksel uji). Baseline *dark-spot thresholding* (blur 7×7, "
  "intensitas < 70; ambang dipilih pada data latih) disertakan, dan generalisasi lintas sensor diuji pada 776 citra PALSAR.")

H2("2.3 Edge AI, poligon, dan simulasi respons")
P("Model DL diekspor ke ONNX lalu dikuantisasi INT8 statis (QDQ per-kanal) [12] dengan kalibrasi 200 citra latih; varian web adalah varian "
  "terkecil yang penurunan Dice validasinya ≤ 0,005. Model ML diekspor sebagai satu graf ONNX (graf fitur + *ensemble* pohon). Aplikasi "
  "web statis di Vercel menjalankan kelima model dengan onnxruntime-web [14] (WebAssembly multi-thread), sehingga citra pengguna tidak "
  "pernah dikirim ke server. Setelah inferensi, mask dipecah menjadi komponen terhubung, batas luar tiap slick ditelusuri lalu "
  "disederhanakan dengan Douglas–Peucker, sehingga diperoleh luas, keliling, dan poligon GeoJSON per slick.")
P("Simulasi skenario (*what-if*) memproyeksikan poligon ke depan. **Drift**: kecepatan slick = 3% kecepatan angin + 100% arus permukaan, "
  "aturan praktis yang juga dipakai model trajektori NOAA GNOME [19]. **Penyebaran**: persamaan Fay [20] rezim gravitasi-viskos "
  "A ∝ ∛(Δ·g·V²/√ν)·√t, dengan volume V = luas × ketebalan menurut kode penampakan Bonn Agreement [21], ditambah difusi turbulen "
  "ΔA = 8πK·t dengan K dari hukum skala Okubo [22] (K = 0,0103 × L pangkat 1,15, dalam cm²/s, L = diameter slick dalam cm). **Kebutuhan oil boom** = 1,3 × keliling *convex hull* slick "
  "saat tim tiba, dengan peringatan bila kecepatan relatif > 0,7 knot karena boom penahan mulai bocor [23].")

# ============================ HASIL DAN PEMBAHASAN ============================
H1("3. Hasil dan Pembahasan")
H2("3.1 Perbandingan model pada set uji")
table("Tabel 2. Hasil pada set uji Sentinel-1 (839 citra, agregasi global) dan uji lintas sensor PALSAR (776 citra).",
      ["Model", "Thr", "Dice/F1", "IoU", "Precision", "Recall", "PR-AUC", "Dice PALSAR"],
      [["Baseline dark-spot", "–", "0,7259", "0,5698", "0,6669", "0,7964", "–", "0,6410"],
       ["ML: Random Forest", "0,40", "0,8036", "0,6717", "0,7721", "0,8377", "0,8981", "0,7506"],
       ["ML: LightGBM", "0,35", "0,8056", "0,6745", "0,7568", "0,8612", "0,9020", "0,7490"],
       ["DL: U-Net Lite", "0,45", "0,8547", "0,7462", "0,8058", "**0,9099**", "0,9361", "0,7774"],
       ["**DL: DeepLabV3+ MobileNetV2**", "0,50", "**0,8667**", "**0,7648**", "**0,8344**", "0,9016", "**0,9479**", "**0,7975**"]],
      [4.6, 1.0, 1.6, 1.4, 1.6, 1.5, 1.5, 1.8], size=8.5, bold_rows=(4,),
      note="Thr = threshold yang dipilih pada data validasi. Baseline mask kosong mencapai akurasi piksel 65,3% tanpa mendeteksi minyak "
           "sama sekali, sehingga akurasi tidak dipakai sebagai metrik utama. Tanpa 9 citra near-duplicate, Dice U-Net Lite hanya turun 0,0008.")
P("Kedua model DL mengungguli model ML sekitar 5–6 poin Dice dan 7–9 poin IoU, dan semua model jauh di atas baseline (Tabel 2, "
  "Gambar 2). Selisih terbesar ada pada **precision**: fitur lokal model ML menandai hampir semua area gelap, termasuk tepi slick dan "
  "bercak *look-alike*, sedangkan DL memanfaatkan konteks spasial yang lebih luas. Model ML tetap 8 poin di atas baseline dengan ukuran "
  "yang sangat kecil. *Feature importance* menunjukkan fitur terpenting adalah minimum lokal 11 piksel, rata-rata lokal 11 piksel, dan "
  "selisih terhadap rata-rata citra, yaitu ukuran kegelapan relatif yang memang menjadi ciri slick pada SAR [1].")
P("DeepLabV3+ pralatih unggul atas U-Net Lite pada Dice (+1,2 poin), precision (+2,9 poin), dan generalisasi lintas sensor (+2,0 poin), "
  "sedangkan U-Net Lite sedikit lebih tinggi recall-nya. Hal ini menunjukkan manfaat *transfer learning* meski citra SAR berbeda domain "
  "dengan ImageNet. Recall semua model DL ≈ 0,90–0,91; untuk peringatan dini HSSE sifat ini lebih aman karena tumpahan yang terlewat lebih "
  "mahal dampaknya daripada alarm palsu yang dapat diverifikasi analis.")
figure(FIG / "model_predictions_side_by_side.png", 15.5, "Gambar 2. Prediksi kelima model pada citra uji mudah, tipikal, sulit, dan tanpa minyak (merah = minyak).")

H2("3.2 Perbandingan dengan tolok ukur dataset SOS")
table("Tabel 3. Perbandingan pada subset uji Sentinel-1 dataset SOS.",
      ["Model", "mIoU (%)", "F1 (%)", "Recall (%)", "Precision (%)", "Ukuran model"],
      [["U-Net [3]", "81,46", "86,10", "81,22", "85,61", "tidak dilaporkan"],
       ["DeepLabv3 [3]", "82,94", "87,70", "84,76", "88,08", "tidak dilaporkan"],
       ["CBD-Net [3]", "83,42", "87,87", "87,32", "91,20", "tidak dilaporkan"],
       ["LightGBM (ours)", "73,43", "80,56", "86,12", "75,68", "0,24 MB"],
       ["U-Net Lite (ours, INT8)", "79,46", "85,45", "90,90", "80,62", "2,01 MB"],
       ["**DeepLabV3+ MobileNetV2 (ours, INT8)**", "**81,23**", "**86,73**", "**90,16**", "**83,44**", "**4,84 MB**"]],
      [5.0, 1.8, 1.6, 1.8, 2.0, 2.6], size=8.5, bold_rows=(5,),
      note="Nilai [3] dikutip dari Tabel III paper tersebut (threshold 0,5; cara agregasi tidak dirinci). mIoU kami = rata-rata IoU minyak dan "
           "IoU latar; recall/precision DeepLabV3+ dari model PyTorch, F1 dari varian INT8. Perbedaan protokol membuat perbandingan ini indikatif.")
P("DeepLabV3+ mempersempit selisih F1 dengan CBD-Net [3] dari 2,4 poin (U-Net Lite) menjadi 1,1 poin, dengan model INT8 4,84 MB yang berjalan "
  "di browser. Recall model DL kami (90–91%) adalah yang tertinggi pada Tabel 3.")

H2("3.3 Efisiensi komputasi (Green/Edge AI)")
table("Tabel 4. Ukuran, kesetaraan ONNX, dan latensi model yang dipakai di web.",
      ["Model (varian web)", "Ukuran", "Dice uji ONNX", "Kesetaraan dengan notebook", "Latensi CPU*", "Latensi browser**"],
      [["Dark-spot (ONNX)", "<0,1 MB", "0,7259", "piksel identik 100%", "0,3 ms", "±1 ms"],
       ["Random Forest (fitur + pohon)", "1,83 MB", "–", "mask sama 99,98%†", "218 ms", "±270 ms"],
       ["LightGBM (fitur + pohon)", "0,24 MB", "–", "mask sama 99,98%†", "349 ms", "±410 ms"],
       ["U-Net Lite INT8", "2,01 MB", "0,8545", "FP32 7,77 MB: 0,8547", "50 ms", "105–122 ms"],
       ["**DeepLabV3+ INT8**", "**4,84 MB**", "**0,8673**", "FP32 17,5 MB: 0,8667", "**26 ms**", "±130 ms"]],
      [4.2, 1.7, 2.0, 3.6, 1.9, 2.3], size=8.5, bold_rows=(4,),
      note="*onnxruntime CPU Colab (2 vCPU), 1 thread, batch 1, median 20 run. **Chrome di laptop uji, onnxruntime-web WASM multi-thread "
           "(COOP/COEP), median 3 run setelah warm-up. †Diuji pada 100 citra uji acak karena inferensi pohon di CPU lambat.")
P("Kuantisasi INT8 memperkecil DeepLabV3+ 3,6 kali tanpa penurunan Dice, dan justru tercepat di CPU karena MobileNetV2 memakai konvolusi "
  "*depthwise*. Model ML berukuran paling kecil, tetapi paling lambat karena setiap piksel (65.536 per citra) harus melewati semua pohon. "
  "Semua model jauh di bawah target panduan (20–30 MB).")

H2("3.4 Analisis error")
P("Dice per citra U-Net Lite memiliki median 0,797 (persentil-10 0,517; persentil-90 0,965). Kesalahan utama semua model adalah: (1) "
  "tumpahan sangat kecil dan garis gelap tipis yang tidak dilabeli (kemungkinan *look-alike* atau label terlewat); (2) batas slick, karena "
  "label berupa poligon kasar sedangkan prediksi lebih halus dan sedikit melebar; dan (3) pada model ML, bercak gelap kecil di area laut "
  "(Gambar 2). Performa per citra juga bervariasi: pada beberapa citra sulit, model ML dapat lebih baik daripada DL, sehingga aplikasi web "
  "menyediakan mode \"bandingkan semua model\" untuk analis. Pada PALSAR, semua model turun 5–9 poin, sehingga *fine-tuning* diperlukan "
  "sebelum dipakai pada sensor atau wilayah lain.")

# ========================= DAMPAK OPERASIONAL & BISNIS ========================
H1("4. Dampak Operasional, Efisiensi Biaya, dan Manajemen Risiko")
figure(FIG / "web_v2_demo.jpg", 14.5, "Gambar 3. Aplikasi web: pilihan 5 model, area terdeteksi dengan poligon bernomor per slick, citra asli, dan peta probabilitas.")
P("Solusi ini menyasar **Pilar HSSE** dan **efisiensi biaya**. Skenario penggunaannya adalah *screening* citra Sentinel-1 di sekitar anjungan, "
  "jalur pipa, dan terminal. Keluarannya (Gambar 3): poligon setiap slick dengan luas (km²) dan keliling, GeoJSON untuk GIS, dan simulasi "
  "respons (Gambar 4). Sebagai contoh, pada citra uji tipikal dengan resolusi 10 m/piksel, DeepLabV3+ mendeteksi 1,145 km² minyak. Dengan "
  "skenario angin 8 m/s dari timur laut, arus 0,3 m/s ke timur, dan ketebalan *rainbow* 5 µm, slick bergerak 0,21 m/s ke tenggara, "
  "melebar ±0,10 km²/jam, dan saat tim tiba 3 jam kemudian telah bergeser 2,3 km dengan luas ±1,45 km². Kebutuhan *oil boom* untuk "
  "mengurung seluruh slick sekitar 16,4 km.")
figure(FIG / "web_v2_sim.jpg", 12.0, "Gambar 4. Simulasi drift & penyebaran: posisi slick +6 jam (warna pelangi), jejak waktu (garis putus), dan oil boom saat tim tiba (pelampung oranye).")
P("**Efisiensi biaya.** Kalkulator di web membandingkan patroli pencarian konvensional dengan CV + verifikasi analis. Dengan asumsi "
  "ilustratif (12 sortie/bulan × 4 jam × Rp45 juta/jam; 10 citra/bulan × 0,5 jam analis × Rp150 ribu/jam; 25% sortie tetap dibutuhkan "
  "untuk konfirmasi), biaya pemantauan turun dari Rp2,16 miliar menjadi Rp0,54 miliar per bulan (−75%). Citra Sentinel-1 tersedia gratis "
  "melalui Copernicus dan inferensi berjalan di browser sehingga biaya komputasinya praktis nol. Angka ini bergantung pada asumsi dan "
  "harus diganti dengan data kontrak perusahaan sebelum diklaim.")
table("Tabel 5. Register risiko HSSE kegiatan pemantauan (L = kemungkinan, S = keparahan, skala 1–5).",
      ["Bahaya", "S", "Sebelum (L×S)", "Sesudah CV (L×S)"],
      [["Paparan uap hidrokarbon (VOC, H₂S) saat inspeksi jarak dekat", "4", "3×4 = 12", "1×4 = 4"],
       ["Kecelakaan penerbangan pengintaian", "5", "2×5 = 10", "1×5 = 5"],
       ["Insiden kapal patroli (jatuh ke laut, tabrakan)", "4", "3×4 = 12", "1×4 = 4"],
       ["Kebakaran/ledakan di dekat minyak segar", "5", "2×5 = 10", "1×5 = 5"],
       ["Kelelahan kru akibat patroli panjang & malam", "3", "4×3 = 12", "2×3 = 6"],
       ["Tumpahan terlambat diketahui (malam/berawan)", "4", "3×4 = 12", "2×4 = 8"],
       ["*Risiko baru:* salah deteksi / slick terlewat", "3", "–", "3×3 = 9"],
       ["**Total skor risiko**", "", "**68**", "**41 (−40%)**"]],
      [8.6, 0.8, 3.0, 3.4], size=8.5)
P("**Manajemen risiko.** Pengurangan sortie pencarian menurunkan jam paparan personel di lapangan dari 192 menjadi 48 jam per bulan "
  "(−75%), sehingga kemungkinan (L) bahaya berbasis paparan turun dua tingkat, sedangkan keparahan (S) tetap karena konsekuensinya sama bila "
  "kejadian terjadi (Tabel 5). Pada hierarki pengendalian, pemindahan pencarian slick dari lapangan ke analisis citra termasuk "
  "**eliminasi paparan**. CV tidak menggantikan tim lapangan sepenuhnya: verifikasi, pengambilan sampel, dan pemasangan boom tetap "
  "dilakukan manusia, dan risiko baru berupa salah deteksi dikendalikan dengan verifikasi analis untuk setiap alarm.")
P("**Batasan.** Dataset SOS tidak memiliki koordinat dan waktu akuisisi sehingga simulasi tidak dapat divalidasi terhadap kejadian nyata; "
  "simulasi juga tidak memodelkan penguapan, emulsifikasi, garis pantai, maupun angin/arus yang berubah. Simulasi diposisikan sebagai "
  "gambaran awal yang diperbarui dengan data cuaca/oseanografi dan citra berikutnya, seperti pada sistem operasional Petalas dkk. [10].")

# ================================ KESIMPULAN ==================================
H1("5. Kesimpulan dan Saran")
P("Pada protokol yang sama, model DL mengungguli model ML untuk segmentasi tumpahan minyak Sentinel-1: DeepLabV3+ MobileNetV2 pralatih "
  "mencapai Dice 0,867 dan IoU 0,765 (1,1 poin F1 di bawah CBD-Net), U-Net Lite 0,855, sedangkan Random Forest dan LightGBM sekitar 0,805 "
  "dengan ukuran model jauh lebih kecil. Kelima model berjalan di browser, dan aplikasi web melengkapi deteksi dengan poligon per slick, "
  "simulasi arah dan laju penyebaran, estimasi kebutuhan oil boom, serta analisis biaya dan risiko HSSE.")
P("Saran pengembangan: (1) menambah kelas *look-alike* dan data multi-sensor; (2) memproses citra Sentinel-1 penuh dengan *tiling* dan "
  "georeferensi sehingga poligon memiliki koordinat nyata; (3) menghubungkan simulasi dengan data angin dan arus operasional (misalnya "
  "Copernicus Marine) serta memvalidasinya dengan citra berurutan; (4) melakukan uji beberapa *seed* untuk mengukur variasi hasil; dan (5) "
  "validasi pilot di perairan Indonesia bersama tim HSSE untuk mengganti asumsi biaya dengan data nyata.")

# =============================== DAFTAR PUSTAKA ===============================
H1("Daftar Pustaka")
refs = [
    "M. Fingas and C. Brown, \"A review of oil spill remote sensing,\" *Sensors*, vol. 18, no. 1, p. 91, 2018, doi: 10.3390/s18010091.",
    "R. Al-Ruzouq *et al.*, \"Sensors, features, and machine learning for oil spill detection and monitoring: A review,\" *Remote Sensing*, vol. 12, no. 20, p. 3338, 2020, doi: 10.3390/rs12203338.",
    "Q. Zhu *et al.*, \"Oil spill contextual and boundary-supervised detection network based on marine SAR images,\" *IEEE Trans. Geosci. Remote Sens.*, vol. 60, pp. 1–10, 2022, doi: 10.1109/TGRS.2021.3115492.",
    "M. Krestenitis *et al.*, \"Oil spill identification from satellite images using deep neural networks,\" *Remote Sensing*, vol. 11, no. 15, p. 1762, 2019, doi: 10.3390/rs11151762.",
    "M. Shaban *et al.*, \"A deep-learning framework for the detection of oil spills from SAR data,\" *Sensors*, vol. 21, no. 7, p. 2351, 2021, doi: 10.3390/s21072351.",
    "R. Hasimoto-Beltran, M. Canul-Ku, G. M. Díaz Méndez, F. J. Ocampo-Torres, and B. Esquivel-Trava, \"Ocean oil spill detection from SAR images based on multi-channel deep learning semantic segmentation,\" *Marine Pollution Bulletin*, vol. 188, p. 114651, 2023, doi: 10.1016/j.marpolbul.2023.114651.",
    "C. Li, D. Kim, S. Park, J. Kim, and J. Song, \"A self-evolving deep learning algorithm for automatic oil spill detection in Sentinel-1 SAR images,\" *Remote Sensing of Environment*, vol. 299, p. 113872, 2023, doi: 10.1016/j.rse.2023.113872.",
    "L. Liao, Q. Zhao, and W. Song, \"Monitoring of oil spill risk in coastal areas based on polarimetric SAR satellite images and deep learning theory,\" *Sustainability*, vol. 15, no. 19, p. 14504, 2023, doi: 10.3390/su151914504.",
    "K. Das, P. Janardhan, and M. R. Singh, \"Oil spill detection in SAR images: A U-Net semantic segmentation framework with multiple backbones,\" in *Water and Environment, Volume 2*, Lecture Notes in Civil Engineering, vol. 414. Singapore: Springer, 2024, pp. 65–77, doi: 10.1007/978-981-97-7502-6_6.",
    "S. Petalas, T. Psomouli, N. Kokkos, and G. Sylaios, \"Operational oil-spill detection in the framework of digital twin using SAR imagery and deep learning,\" in *Proc. OCEANS 2025 Brest*, 2025, pp. 1–6, doi: 10.1109/OCEANS58557.2025.11104698.",
    "O. Ronneberger, P. Fischer, and T. Brox, \"U-Net: Convolutional networks for biomedical image segmentation,\" in *Proc. MICCAI 2015*, LNCS, vol. 9351. Cham: Springer, 2015, pp. 234–241, doi: 10.1007/978-3-319-24574-4_28.",
    "B. Jacob *et al.*, \"Quantization and training of neural networks for efficient integer-arithmetic-only inference,\" in *Proc. IEEE/CVF CVPR*, 2018, pp. 2704–2713, doi: 10.1109/CVPR.2018.00286.",
    "BitsandLayers, \"Deep-SAR SOS Oil Spill Detection Dataset,\" Kaggle, v1, 2026. [Online]. Available: https://www.kaggle.com/datasets/bitsandlayers/sar-oil-spill-segmentation-dataset-sos (diakses 24 Sep. 2026). Lisensi CC BY 4.0.",
    "ONNX Runtime developers, \"ONNX Runtime (onnxruntime-web) v1.30.0,\" 2026. [Online]. Available: https://onnxruntime.ai",
    "L. Breiman, \"Random forests,\" *Machine Learning*, vol. 45, no. 1, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.",
    "G. Ke *et al.*, \"LightGBM: A highly efficient gradient boosting decision tree,\" in *Advances in Neural Information Processing Systems 30 (NIPS 2017)*, 2017, pp. 3146–3154.",
    "L.-C. Chen, Y. Zhu, G. Papandreou, F. Schroff, and H. Adam, \"Encoder-decoder with atrous separable convolution for semantic image segmentation,\" in *Proc. ECCV 2018*, LNCS, vol. 11211. Cham: Springer, 2018, pp. 833–851, doi: 10.1007/978-3-030-01234-2_49.",
    "M. Sandler, A. Howard, M. Zhu, A. Zhmoginov, and L.-C. Chen, \"MobileNetV2: Inverted residuals and linear bottlenecks,\" in *Proc. IEEE/CVF CVPR*, 2018, pp. 4510–4520, doi: 10.1109/CVPR.2018.00474.",
    "NOAA Office of Response and Restoration, \"GNOME: General NOAA Operational Modeling Environment,\" 2024. [Online]. Available: https://response.restoration.noaa.gov/gnome",
    "J. A. Fay, \"Physical processes in the spread of oil on a water surface,\" in *Proc. Joint Conf. Prevention and Control of Oil Spills*, Washington, DC: API, 1971, pp. 463–467, doi: 10.7901/2169-3358-1971-1-463.",
    "Bonn Agreement, \"Bonn Agreement Aerial Operations Handbook: Bonn Agreement Oil Appearance Code,\" 2016. [Online]. Available: https://www.bonnagreement.org",
    "A. Okubo, \"Oceanic diffusion diagrams,\" *Deep Sea Research and Oceanographic Abstracts*, vol. 18, no. 8, pp. 789–802, 1971, doi: 10.1016/0011-7471(71)90046-5.",
    "ITOPF, \"Use of booms in oil pollution response,\" Technical Information Paper no. 3, London: ITOPF, 2011. [Online]. Available: https://www.itopf.org",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"[{i}]\t"); r.font.size = Pt(9.5)
    add_runs(p, ref, size=9.5)

# ================================= LAMPIRAN ===================================
from docx.enum.text import WD_BREAK
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
H1("Lampiran A. Contribution Statement")
P("Tabel A1 merinci peran, tanggung jawab, dan persentase kontribusi tiap anggota (total 100%). Ketua mendapat porsi lebih besar "
  "karena merangkap koordinasi tim dan pemodelan inti.", indent=False)
table("Tabel A1. Pembagian peran dan kontribusi anggota Kelompok 1 CAYMAN.",
      ["Nama (NPM)", "Peran", "Tanggung jawab", "Kontribusi"],
      [["Ade Rizky Darmawan (23083010080), Ketua", "Model Architect Specialist", "Koordinasi tim; perancangan arsitektur U-Net Lite, training pipeline di GPU, pemilihan loss dan threshold, evaluasi metrik", "20%"],
       ["Arkananta Daniswara Handoyo (23083010059)", "Data & Pipeline Specialist", "Akuisisi dataset SOS, audit data (pasangan, label, duplikat), preprocessing, pembagian train/val/test", "16%"],
       ["Muhammad Arsyad Alzam (23083010082)", "Deployment & Edge Specialist", "Ekspor ONNX, kuantisasi FP16/INT8, benchmark latensi, onnxruntime-web, deployment Vercel, manajemen GitHub", "16%"],
       ["Choirul Amin (22083010050)", "Frontend & UX Specialist", "Desain dan pengembangan UI web, visualisasi hasil segmentasi, fitur Try Sample Data dan estimasi km², landing page tim", "16%"],
       ["Hana Titania Sastrian (23083010056)", "Lead Technical Writer", "Penyusunan laporan ilmiah, penelusuran dan pengelolaan referensi, dokumentasi README", "16%"],
       ["Zaydan Arief Athallah (23083010063)", "Research & Business Impact Analyst", "Kajian state of the art, analisis dampak operasional HSSE dan estimasi bisnis, materi presentasi", "16%"],
       ["**Total**", "", "", "**100%**"]],
      [4.4, 3.2, 6.4, 1.8], size=8.5)

H1("Lampiran B. Artefak Proyek")
B("Repositori GitHub (publik): https://github.com/aderizkydarmawan/cayman-oil-spill-sar, berisi notebook Colab, skrip model/ekspor, "
  "hasil audit, metrik (JSON/CSV), model ONNX, dan kode web.")
B("Aplikasi web: https://cayman-kelompok1.vercel.app, yang memuat demo inferensi 5 model (*Try Sample Data*), poligon & GeoJSON per slick, simulasi penyebaran & oil boom, kalkulator biaya & risiko HSSE, dan profil tim.")
B("Reproduksi: buka notebooks/Kasus38_SAR_OilSpill_UNetLite.ipynb di Google Colab (runtime T4 GPU) lalu jalankan *Run all*; perbandingan ML vs DL ada di Bagian B (sel 14–19), hasilnya di outputs/v2/.")

doc.save(OUT)
print("saved", OUT)
