"""Tahap 11: laporan gabungan berbahasa awam untuk praktisi (docx + PDF via Word).
Pakai (dari folder uji-gee-balikpapan): PYTHONPATH=scripts python scripts/11_laporan_word.py
Isi mengikuti telaah council 17 Okt (oil-spill-hsse-kelayakan/council-2026-10-17-telaah-laporan.md).
"""
import json, os
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from situs import _S

PENYUSUN = "Ade Rizky Darmawan · wewamazing41@gmail.com"
TANGGAL = "17 Oktober 2026"

# urutan: hari biasa dulu (yang paling sering akan datang), lalu dua kejadian yang diketahui
CONTOH = [
    ("karawang", "2019-01-19", "Contoh 1 — hari biasa (tanpa kejadian yang tercatat)",
     "Tidak ada tumpahan yang tercatat pada tanggal ini. Beginilah isi laporan pada kebanyakan lintasan: "
     "beberapa bercak kecil yang sebagian besar kemungkinan bukan tumpahan.", {}),
    ("karawang", "2019-07-30", "Contoh 2 — kejadian yang diketahui: Karawang",
     "Kebocoran sumur YYA-1 (blok ONWJ) mulai 12 Juli 2019. Tanggal 30 Juli dipilih sebagai contoh saat sebaran sudah luas, "
     "bukan citra pertama setelah kejadian.",
     {5: "±65 km di barat lokasi utama; kemungkinan tidak terkait kejadian ini",
      8: "±65 km di barat lokasi utama; kemungkinan tidak terkait kejadian ini"}),
    ("balikpapan", "2018-04-01", "Contoh 3 — kejadian yang diketahui: Teluk Balikpapan",
     "Pipa minyak bawah laut di Teluk Balikpapan putus pada 31 Maret 2018. Citra ini diambil sekitar sehari setelahnya.",
     {2: "di alur sungai, menempel daratan; kemungkinan bukan minyak atau batas darat kurang rapi"}),
]

c = lambda x, nd=1: f"{x:.{nd}f}".replace(".", ",")

doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("PROTOTIPE – DATA ARSIP"); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Prototipe riset ADE, bukan produk resmi. Berisi bercak gelap di citra radar, bukan konfirmasi tumpahan, "
               "dan tidak menyatakan sumber atau pihak yang bertanggung jawab. Mohon tidak disebarkan tanpa konteks.")
r.font.size = Pt(8); r.italic = True

def para(text, bold=False, italic=False, size=None, style=None):
    p = doc.add_paragraph(style=style); rr = p.add_run(text); rr.bold = bold; rr.italic = italic
    if size: rr.font.size = Pt(size)
    return p
bullets = lambda xs, style="List Bullet": [para(x, style=style) for x in xs]

doc.add_heading("Contoh Laporan Bercak Gelap di Laut dari Citra Radar Satelit", 0)
para(f"Disusun oleh {PENYUSUN} · {TANGGAL}", size=9.5)

doc.add_heading("Intinya (30 detik)", 1)
bullets(["Setiap kali satelit radar Sentinel-1 melintas (±6–12 hari sekali), laporan seperti ini bisa dibuat otomatis: "
         "peta dan daftar koordinat bercak gelap di laut.",
         "Bercak gelap belum tentu minyak. Pada hari biasa tetap muncul beberapa bercak (median 3–6,5 per lintasan).",
         "Laporan ini bukan peringatan dini. Kegunaan yang paling mungkin: pemantauan rutin di area yang jarang dipatroli, "
         "dan pemetaan sebaran setelah kejadian.",
         "Ini uji sekali jalan dari data arsip, bukan layanan yang sedang berjalan."])

doc.add_heading("Tiga pertanyaan untuk praktisi", 1)
bullets(["Informasi apa yang harus ada supaya praktisi mengirim kapal atau pesawat untuk mengecek, dan apa yang membuat laporan seperti ini diabaikan?",
         "Kalau setiap laporan berisi 3–7 bercak seperti Contoh 1, berapa bercak per minggu yang masih sanggup dicek?",
         "Apakah saat ini sudah ada layanan atau sistem serupa yang dipakai? Kalau ada, apa yang masih kurang?"], "List Number")
para("Jawaban singkat lewat pesan atau telepon 15 menit sudah sangat membantu. Jawaban hanya dipakai untuk riset ini, "
     "tidak akan dikutip atas nama praktisi tanpa izin.", italic=True, size=9.5)

doc.add_heading("Cara membaca laporan", 1)
bullets(["Peta: abu-abu = laut di radar (makin gelap makin licin), krem = darat, merah = bercak gelap yang ditandai, dengan nomor bercak.",
         "Lokasi: bercak yang berjarak kurang dari 5 km digabung menjadi satu lokasi.",
         "Angin lemah (< 3 m/s): laut yang tenang memang tampak gelap, jadi bercak di area itu belum bisa dinilai. "
         "Angin diukur langsung dari citra di sekitar tiap bercak.",
         "Bercak yang menempel garis pantai sering berasal dari muara sungai atau batas darat yang kurang rapi.",
         "Koordinat ditulis lintang, bujur (pakai titik) dan bisa ditempel langsung ke Google Maps.",
         "Lokasi sumber kejadian tidak ditandai di peta karena kami tidak memiliki koordinat resminya."])

for name, tgl, judul, konteks, ragu in CONTOH:
    S = _S[name]; d = os.path.join(S["results"], f"laporan_{tgl}")
    a = pd.read_csv(os.path.join(d, f"alarm_{tgl}.csv")); meta = json.load(open(os.path.join(d, "meta.json")))
    doc.add_page_break()
    doc.add_heading(judul, 1)
    para(f"{name.title()} · satelit lewat {meta['waktu_lokal']}", bold=True)
    para(konteks, italic=True)

    tot = a.groupby("kelompok").luas_km2.sum().sort_values(ascending=False)
    gb = a[a.kelompok == tot.index[0]]; n_low = int((a.angin_sekitar_ms < 3).sum())
    txt = (f"Ada {len(a)} bercak gelap dengan total luas sekitar {c(a.luas_km2.sum())} km², di {len(tot)} lokasi. "
           f"Lokasi terbesar ada di sekitar {gb.lat.mean():.3f}, {gb.lon.mean():.3f} ({len(gb)} bercak, {c(gb.luas_km2.sum())} km²). ")
    txt += (f"{n_low} bercak berada di area berangin lemah sehingga belum bisa dinilai."
            if n_low else f"Angin di sekitar semua bercak cukup ({c(a.angin_sekitar_ms.min())}–{c(a.angin_sekitar_ms.max())} m/s), "
                          "jadi semua bercak bisa dinilai; pemeriksaan angin tidak menyisihkan bercak mana pun di contoh ini.")
    para(txt)
    if name == "balikpapan":
        para("Model cuaca global (ERA5) mencatat angin umum yang lemah pagi itu, tetapi angin yang diukur dari citra di sekitar bercak "
             "3–6 m/s. Angin dari citra dipakai karena diukur tepat di lokasi dan waktu yang sama, sedangkan ERA5 berupa rata-rata "
             "kasar (kotak ±30 km per jam).", size=9.5)

    doc.add_picture(os.path.join(d, f"laporan_{tgl}.png"), width=Cm(17))

    t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
    for cc, h in zip(t.rows[0].cells, ["Lokasi", "Bercak no.", "Luas", "Koordinat pusat", "Catatan"]):
        cc.text = h
    for i, k in enumerate(tot.index, 1):
        g = a[a.kelompok == k]; row = t.add_row().cells
        notes = [ragu[j] for j in g.id if j in ragu]
        if (g.angin_sekitar_ms < 3).any(): notes.append("angin lemah, belum bisa dinilai")
        if (g.jarak_pantai_km < 1).any() and not notes: notes.append("menempel pantai, cek muara sungai")
        row[0].text = str(i); row[1].text = ", ".join(map(str, g.id)); row[2].text = f"{c(g.luas_km2.sum())} km²"
        row[3].text = f"{g.lat.mean():.4f}, {g.lon.mean():.4f}"; row[4].text = "; ".join(dict.fromkeys(notes)) or "-"
    for row in t.rows:
        for cc in row.cells:
            for p in cc.paragraphs:
                for rr in p.runs: rr.font.size = Pt(9)

doc.add_page_break()
doc.add_heading("Yang sebaiknya dilakukan saat laporan datang", 1)
bullets(["Cocokkan lokasi dengan aset terdekat (anjungan, pipa, jalur kapal) dan laporan lapangan hari itu.",
         "Bila memungkinkan, periksa dengan pengamatan udara, kapal patroli, atau citra optik (Sentinel-2) pada hari yang sama atau sesudahnya.",
         "Dahulukan lokasi terbesar dan yang paling dekat aset. Bercak yang menempel garis pantai sering berasal dari muara sungai; cek dulu."])

doc.add_heading("Seberapa sering alat ini keliru?", 1)
para("Alat ini dijalankan pada citra hari-hari tanpa tumpahan yang tercatat: 82 citra di Teluk Balikpapan dan 77 citra di pesisir "
     "Karawang (2017–2019). Pada hari-hari itu tetap muncul median 3–6,5 bercak per lintasan. Penyebab utamanya angin lemah. "
     "Sebagian bercak itu bisa jadi buangan kapal yang nyata, tetapi tidak ada data lapangan untuk memastikannya.")

doc.add_heading("Yang perlu diingat", 1)
bullets(["Ini dugaan dari citra radar, bukan kepastian, dan belum ada konfirmasi lapangan.",
         "Sudah dijalankan pada 14 kejadian arsip di Indonesia dan sekitarnya. Lapisan minyak di laut terbuka tertangkap jelas pada 3 "
         "kejadian (Balikpapan 2018, Karawang 2019, Karawang 2021) dan mungkin tertangkap pada 2 kejadian (Kep. Seribu 2020, Singapura 2024). "
         "Gumpalan minyak (sludge) di pantai Batam, Bintan, dan Lampung tidak tertangkap. Pengaturan skala kecerahan dipilih setelah "
         "melihat kejadian Balikpapan, jadi uji ini belum sepenuhnya independen.",
         "Kemampuan menemukan tumpahan sudah diuji pada data berlabel ahli dari dua sumber terpisah: 73% (Mediterania, 2.630 tumpahan) "
         "dan 81% (uji independen Trujillo, 670 tumpahan); 92% piksel yang ditandai memang minyak.",
         "Alat belum bisa membedakan minyak dari bercak alami yang mirip dengan andal. Model pembeda gagal pada uji independen "
         "(hanya 76% luas minyak diberi prioritas tinggi, target 90%), jadi hanya dipakai untuk mengurutkan; tidak ada bercak yang dibuang. "
         "Untuk menyetelnya dibutuhkan penilaian praktisi atas contoh bercak perairan Indonesia.",
         "Tumpahan bisa terlewat. Di Karawang yang terdeteksi ±11 km², sedangkan perkiraan LEMIGAS ±25 km²; bagian minyak yang tipis "
         "tidak selalu tampak. Tumpahan juga bisa terlewat bila terjadi di antara dua lintasan atau di area berangin lemah.",
         "Citra pertama baru tersedia ±1,5 hari setelah kejadian, lalu tiap 6–12 hari. Jadi ini bukan peringatan dini.",
         "Bercak gelap juga bisa berasal dari lapisan alami (alga, sisa hujan, buangan sungai).",
         "Tahap kesiapan: purwarupa yang diuji dengan data arsip, belum diuji di lapangan, dan belum ada yang menjalankannya secara rutin."])

doc.add_heading("Lampiran: cara kerja singkat", 1)
para("Citra radar Sentinel-1 diambil dari Google Earth Engine. Model pengenal pola (U-Net, dari riset sebelumnya, tanpa dilatih ulang) "
     "menandai bercak gelap di laut, dan bercak lebih kecil dari 0,1 km² dibuang. Angin di sekitar tiap bercak dihitung dari citra itu "
     "sendiri (metode CMOD5.N, cincin 1–3 km di sekeliling bercak). Semua langkah berjalan di laptop dengan data gratis. "
     "Mengandung data Copernicus Sentinel-1 (ESA) dan ERA5 (ECMWF). Kode dan catatan lengkap tersedia bila dibutuhkan.", size=9.5)

os.makedirs("laporan", exist_ok=True)
out = os.path.abspath(os.path.join("laporan", "contoh_laporan_tumpahan.docx")); doc.save(out)
import win32com.client
w = win32com.client.DispatchEx("Word.Application"); w.Visible = False
try:
    dd = w.Documents.Open(out); dd.SaveAs(out[:-5] + ".pdf", FileFormat=17); dd.Close(False)
finally:
    w.Quit()
print("->", out, "+ .pdf")
